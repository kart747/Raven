"""
Apply verified FCRA registration statuses from an official MHA list.

Download a list from fcraonline.nic.in (e.g. "Cancelled associations"), save it as
CSV or XLSX, then:

    python -m app.cli ingest-fcra-status cancelled.xlsx --status Cancelled \
        --source-url "https://fcraonline.nic.in/..."

The registration-number column is detected from its header. Statuses are merged into
data/fcra_status_overrides.csv (so a full FCRA re-import keeps them) and applied to the
NGOs table immediately.
"""
import csv
import os
import re

from .database import SessionLocal, ensure_schema
from . import models

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
OVERRIDES_PATH = os.environ.get(
    "FCRA_STATUS_CSV", os.path.join(PROJECT_ROOT, "data", "fcra_status_overrides.csv")
)
VALID_STATUSES = {"Active", "Suspended", "Cancelled"}
FIELDS = ["fcra_registration_number", "status", "source_url"]


def _read_rows(path: str) -> list[list[str]]:
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        sheet = load_workbook(path, read_only=True, data_only=True).active
        return [["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.reader(f))


def extract_registration_numbers(rows: list[list[str]]) -> list[str]:
    """Find the header row with a registration column and return normalised numbers below it."""
    for i, row in enumerate(rows[:20]):
        for j, cell in enumerate(row):
            if re.search(r"(regist\w*|reg\.?|fcra)\s*(no\b|no\.|number|num\b)", cell, re.I):
                numbers = []
                for data in rows[i + 1:]:
                    if j < len(data):
                        digits = re.sub(r"\D", "", data[j])
                        if len(digits) >= 5:
                            numbers.append(digits)
                return numbers
    raise ValueError("No registration-number column found (expected a header like 'Registration No.').")


def _load_overrides() -> dict[str, dict]:
    if not os.path.exists(OVERRIDES_PATH):
        return {}
    with open(OVERRIDES_PATH, newline="", encoding="utf-8") as f:
        return {r["fcra_registration_number"]: r for r in csv.DictReader(f)}


def run_import(path: str, status: str, source_url: str) -> dict:
    status = status.title()
    if status not in VALID_STATUSES:
        raise ValueError(f"status must be one of {sorted(VALID_STATUSES)}")
    if not source_url:
        raise ValueError("--source-url is required so every status can be traced to MHA")

    numbers = extract_registration_numbers(_read_rows(path))
    overrides = _load_overrides()
    for reg in numbers:
        overrides[reg] = {"fcra_registration_number": reg, "status": status, "source_url": source_url}

    os.makedirs(os.path.dirname(OVERRIDES_PATH), exist_ok=True)
    with open(OVERRIDES_PATH, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(sorted(overrides.values(), key=lambda r: r["fcra_registration_number"]))

    ensure_schema()
    db = SessionLocal()
    try:
        matched = db.query(models.NGO).filter(models.NGO.fcra_registration_number.in_(numbers))\
            .update({models.NGO.registration_status: status}, synchronize_session=False)
        db.commit()
    finally:
        db.close()

    return {
        "numbers_in_file": len(numbers),
        "matched_ngos_in_database": matched,
        "not_in_fcra_returns_data": len(set(numbers)) - matched,
        "overrides_file": OVERRIDES_PATH,
    }
